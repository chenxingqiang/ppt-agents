#!/usr/bin/env python3
"""
PPT Agents - Unified CLI for generating PowerPoint presentations from markdown files.

Usage:
    python main.py generate --input <dir> --output <dir> [--template <path>]
    python main.py template --theme <name> --output <path>
    python main.py template --all
    python main.py check
"""

import argparse
import os
import sys
import logging
import glob
from datetime import datetime
from typing import List, Optional, Dict

# Configure logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)


# =============================================================================
# Environment Check
# =============================================================================

def check_virtual_env() -> bool:
    """Check if running in a virtual environment."""
    in_venv = sys.prefix != sys.base_prefix
    if not in_venv:
        logger.warning("Not running in a virtual environment!")
        logger.info("Please activate your virtual environment:")
        logger.info("  source .venv/bin/activate  # On Unix/Mac")
        logger.info("  .venv\\Scripts\\activate    # On Windows")
        return False
    logger.info(f"Virtual environment active: {sys.prefix}")
    return True


def check_dependencies() -> bool:
    """Check if required dependencies are installed."""
    required = ['pptx', 'markdown', 'bs4', 'yaml', 'frontmatter']
    missing = []
    
    for module in required:
        try:
            __import__(module)
        except ImportError:
            missing.append(module)
    
    if missing:
        logger.error(f"Missing dependencies: {', '.join(missing)}")
        logger.info("Run: pip install -r requirements.txt")
        return False
    
    logger.info("All dependencies installed")
    return True


# =============================================================================
# Template Generation
# =============================================================================

def generate_template(
    theme: str = "modern",
    output_path: str = "templates/modern_template.pptx",
    width: float = 13.33,
    height: float = 7.5
) -> str:
    """
    Generate a PowerPoint template with the specified theme.
    
    Args:
        theme: Theme name (modern, tech, nature, elegant, vibrant, corporate, creative, dark)
        output_path: Output path for the template
        width: Slide width in inches
        height: Slide height in inches
        
    Returns:
        Path to the generated template
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    
    # Theme color definitions
    THEMES = {
        'modern': {
            'primary': RGBColor(0, 82, 147),
            'secondary': RGBColor(0, 114, 198),
            'accent': RGBColor(89, 89, 89),
            'background': RGBColor(255, 255, 255),
            'text': RGBColor(51, 51, 51),
            'font': 'Arial'
        },
        'tech': {
            'primary': RGBColor(28, 28, 35),
            'secondary': RGBColor(75, 207, 250),
            'accent': RGBColor(52, 231, 228),
            'background': RGBColor(18, 18, 24),
            'text': RGBColor(255, 255, 255),
            'font': 'Helvetica'
        },
        'nature': {
            'primary': RGBColor(76, 175, 80),
            'secondary': RGBColor(139, 195, 74),
            'accent': RGBColor(205, 220, 57),
            'background': RGBColor(255, 255, 255),
            'text': RGBColor(51, 51, 51),
            'font': 'Georgia'
        },
        'elegant': {
            'primary': RGBColor(63, 63, 63),
            'secondary': RGBColor(158, 118, 80),
            'accent': RGBColor(213, 196, 161),
            'background': RGBColor(250, 248, 245),
            'text': RGBColor(51, 51, 51),
            'font': 'Garamond'
        },
        'vibrant': {
            'primary': RGBColor(255, 87, 51),
            'secondary': RGBColor(255, 189, 74),
            'accent': RGBColor(102, 187, 255),
            'background': RGBColor(255, 255, 255),
            'text': RGBColor(51, 51, 51),
            'font': 'Verdana'
        },
        'corporate': {
            'primary': RGBColor(25, 51, 92),
            'secondary': RGBColor(52, 152, 219),
            'accent': RGBColor(155, 89, 182),
            'background': RGBColor(255, 255, 255),
            'text': RGBColor(44, 62, 80),
            'font': 'Calibri'
        },
        'creative': {
            'primary': RGBColor(241, 90, 135),
            'secondary': RGBColor(255, 159, 26),
            'accent': RGBColor(88, 177, 159),
            'background': RGBColor(255, 255, 255),
            'text': RGBColor(51, 51, 51),
            'font': 'Century Gothic'
        },
        'dark': {
            'primary': RGBColor(0, 255, 255),
            'secondary': RGBColor(255, 64, 129),
            'accent': RGBColor(83, 83, 83),
            'background': RGBColor(18, 18, 18),
            'text': RGBColor(255, 255, 255),
            'font': 'Consolas'
        }
    }
    
    colors = THEMES.get(theme, THEMES['modern'])
    
    prs = Presentation()
    prs.slide_width = Inches(width)
    prs.slide_height = Inches(height)
    
    # Create slides
    _create_template_slides(prs, colors)
    
    # Save
    os.makedirs(os.path.dirname(output_path) if os.path.dirname(output_path) else '.', exist_ok=True)
    prs.save(output_path)
    logger.info(f"Template saved: {output_path} (theme: {theme})")
    return output_path


def _create_template_slides(prs, colors: Dict):
    """Create all template slide layouts."""
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN
    
    blank_layout = prs.slide_layouts[6]
    
    # 1. Title Slide
    slide = prs.slides.add_slide(blank_layout)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = colors['background']
    
    title = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(12.33), Inches(1.5))
    p = title.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = "Presentation Title"
    run.font.name = colors['font']
    run.font.size = Pt(44)
    run.font.bold = True
    run.font.color.rgb = colors['primary']
    
    subtitle = slide.shapes.add_textbox(Inches(1), Inches(4.2), Inches(11.33), Inches(1))
    p = subtitle.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = "Subtitle"
    run.font.name = colors['font']
    run.font.size = Pt(24)
    run.font.color.rgb = colors['accent']
    
    # 2. Section Slide
    slide = prs.slides.add_slide(blank_layout)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = colors['primary']
    
    title = slide.shapes.add_textbox(Inches(0.5), Inches(3), Inches(12.33), Inches(1.5))
    p = title.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = "Section Title"
    run.font.name = colors['font']
    run.font.size = Pt(40)
    run.font.bold = True
    run.font.color.rgb = colors['background']
    
    # 3. Content Slide
    slide = prs.slides.add_slide(blank_layout)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = colors['background']
    
    title = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.33), Inches(0.8))
    p = title.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = "Slide Title"
    run.font.name = colors['font']
    run.font.size = Pt(32)
    run.font.bold = True
    run.font.color.rgb = colors['primary']
    
    content = slide.shapes.add_textbox(Inches(0.5), Inches(1.3), Inches(12.33), Inches(5.7))
    content.text_frame.word_wrap = True
    p = content.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = "Content"
    run.font.name = colors['font']
    run.font.size = Pt(18)
    run.font.color.rgb = colors['text']
    
    # Footer line
    line = slide.shapes.add_shape(1, Inches(0.5), Inches(7.1), Inches(12.33), Pt(2))
    line.fill.solid()
    line.fill.fore_color.rgb = colors['secondary']
    line.line.fill.background()
    
    # 4. Two-Column Slide
    slide = prs.slides.add_slide(blank_layout)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = colors['background']
    
    title = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.33), Inches(0.8))
    p = title.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = "Two Columns"
    run.font.name = colors['font']
    run.font.size = Pt(32)
    run.font.bold = True
    run.font.color.rgb = colors['primary']
    
    for x in [0.5, 6.83]:
        col = slide.shapes.add_textbox(Inches(x), Inches(1.3), Inches(6), Inches(5.5))
        col.text_frame.word_wrap = True
    
    # 5. Blank Slide
    slide = prs.slides.add_slide(blank_layout)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = colors['background']


def generate_all_templates(output_dir: str = "templates") -> List[str]:
    """Generate all theme templates."""
    themes = ['modern', 'tech', 'nature', 'elegant', 'vibrant', 'corporate', 'creative', 'dark']
    generated = []
    
    for theme in themes:
        output_path = os.path.join(output_dir, f"{theme}_template.pptx")
        generate_template(theme=theme, output_path=output_path)
        generated.append(output_path)
    
    logger.info(f"\nGenerated {len(generated)} templates in {output_dir}/")
    return generated


# =============================================================================
# PPT Generation from Markdown
# =============================================================================

def extract_title_from_markdown(file_path: str) -> str:
    """Extract the title from a markdown file."""
    with open(file_path, 'r', encoding='utf-8') as f:
        content = f.read()
        first_line = content.split('\n')[0]
        title = first_line.strip('# ').replace(':', '-').replace('/', '-')
        return title


def generate_ppt_from_markdown(
    input_dir: str,
    output_dir: str,
    template_path: str = "templates/modern_template.pptx",
    mode: str = "simple",
    clean_existing: bool = False
) -> Dict:
    """
    Generate PPT files from markdown files in a directory.
    
    Args:
        input_dir: Directory containing markdown files
        output_dir: Directory for output PPT files
        template_path: Path to the PowerPoint template
        mode: Generation mode ('simple' or 'detailed')
        clean_existing: Whether to clean existing PPT files first
        
    Returns:
        Dictionary with processing statistics
    """
    stats = {'total': 0, 'successful': 0, 'failed': 0, 'failed_files': []}
    
    os.makedirs(output_dir, exist_ok=True)
    
    # Clean existing files if requested
    if clean_existing:
        old_files = glob.glob(os.path.join(output_dir, "*.pptx"))
        for f in old_files:
            os.remove(f)
            logger.info(f"Cleaned: {f}")
    
    # Get markdown files
    md_files = sorted(glob.glob(os.path.join(input_dir, "*.md")))
    stats['total'] = len(md_files)
    
    if not md_files:
        logger.warning(f"No markdown files found in {input_dir}")
        return stats
    
    logger.info(f"Found {len(md_files)} markdown files")
    
    for md_file in md_files:
        try:
            logger.info(f"Processing: {md_file}")
            
            base_name = os.path.splitext(os.path.basename(md_file))[0]
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            title = extract_title_from_markdown(md_file)
            output_file = os.path.join(output_dir, f"{timestamp}-{base_name}-{title.replace(' ', '')}.pptx")
            
            if mode == "detailed":
                from src.ppt_gen import MarkdownReader, PPTTemplateReader, DetailedPPTGenerator
                template_reader = PPTTemplateReader(template_path)
                markdown_reader = MarkdownReader(md_file)
                generator = DetailedPPTGenerator(template_reader, markdown_reader, output_dir)
                output_file = generator.generate_ppt()
            else:
                from src.ppt_generator import PPTGenerator
                generator = PPTGenerator(template_path)
                generator.process_markdown_file(md_file, output_file)
            
            if os.path.exists(output_file) and os.path.getsize(output_file) > 0:
                logger.info(f"Generated: {output_file}")
                stats['successful'] += 1
            else:
                stats['failed'] += 1
                stats['failed_files'].append(md_file)
                
        except Exception as e:
            logger.error(f"Error processing {md_file}: {e}")
            stats['failed'] += 1
            stats['failed_files'].append(md_file)
    
    # Summary
    logger.info(f"\n{'='*50}")
    logger.info(f"Complete: {stats['successful']}/{stats['total']} successful")
    
    if stats['failed_files']:
        logger.warning(f"Failed ({len(stats['failed_files'])}):")
        for f in stats['failed_files']:
            logger.warning(f"  - {f}")
    
    return stats


# =============================================================================
# CLI Interface
# =============================================================================

def create_parser() -> argparse.ArgumentParser:
    """Create the argument parser."""
    parser = argparse.ArgumentParser(
        prog='ppt-agents',
        description='Generate PowerPoint presentations from markdown files',
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Examples:
  %(prog)s generate -i iechor-bp-md -o iechor-bp-ppt
  %(prog)s generate -i llm-lecture-md -o llm-lecture-ppt --mode detailed
  %(prog)s template --theme modern --output templates/modern.pptx
  %(prog)s template --all
  %(prog)s check
        """
    )
    
    subparsers = parser.add_subparsers(dest='command', help='Available commands')
    
    # Generate command
    gen_parser = subparsers.add_parser('generate', help='Generate PPT from markdown files')
    gen_parser.add_argument('-i', '--input', required=True, help='Input directory with markdown files')
    gen_parser.add_argument('-o', '--output', required=True, help='Output directory for PPT files')
    gen_parser.add_argument('-t', '--template', default='templates/modern_template.pptx', help='Template path')
    gen_parser.add_argument('-m', '--mode', choices=['simple', 'detailed'], default='simple', help='Generation mode')
    gen_parser.add_argument('-c', '--clean', action='store_true', help='Clean existing PPT files first')
    
    # Template command
    tpl_parser = subparsers.add_parser('template', help='Generate PowerPoint templates')
    tpl_parser.add_argument('--theme', choices=['modern', 'tech', 'nature', 'elegant', 'vibrant', 'corporate', 'creative', 'dark'], help='Theme name')
    tpl_parser.add_argument('--output', help='Output path for template')
    tpl_parser.add_argument('--all', action='store_true', help='Generate all theme templates')
    tpl_parser.add_argument('--width', type=float, default=13.33, help='Slide width in inches')
    tpl_parser.add_argument('--height', type=float, default=7.5, help='Slide height in inches')
    
    # Check command
    subparsers.add_parser('check', help='Check environment and dependencies')
    
    return parser


def main():
    """Main entry point."""
    parser = create_parser()
    args = parser.parse_args()
    
    if not args.command:
        parser.print_help()
        return 1
    
    if args.command == 'check':
        venv_ok = check_virtual_env()
        deps_ok = check_dependencies()
        return 0 if (venv_ok and deps_ok) else 1
    
    elif args.command == 'template':
        if args.all:
            generate_all_templates()
        elif args.theme and args.output:
            generate_template(theme=args.theme, output_path=args.output, width=args.width, height=args.height)
        elif args.theme:
            generate_template(theme=args.theme, output_path=f"templates/{args.theme}_template.pptx", width=args.width, height=args.height)
        else:
            # Default: generate modern template
            generate_template()
        return 0
    
    elif args.command == 'generate':
        # Ensure template exists
        if not os.path.exists(args.template):
            logger.warning(f"Template not found: {args.template}")
            logger.info("Generating default template...")
            generate_template(output_path=args.template)
        
        stats = generate_ppt_from_markdown(
            input_dir=args.input,
            output_dir=args.output,
            template_path=args.template,
            mode=args.mode,
            clean_existing=args.clean
        )
        return 0 if stats['failed'] == 0 else 1
    
    return 0


if __name__ == "__main__":
    sys.exit(main())
